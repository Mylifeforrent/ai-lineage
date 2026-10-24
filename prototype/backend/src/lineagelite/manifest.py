"""清单归一化：CSV（双文件）/ XLSX（双 sheet）/ JSON（单文件）→ 统一记录。

对应设计方案 §3.4：三种物理格式归一化为同一 pydantic 模型，后续处理与格式无关。
"""

import csv
import io
import json

from openpyxl import load_workbook

from .models import DatasetFieldRecord, LineageEdgeRecord

DATASET_COLUMNS = ["dataset_id", "dataset_name", "field_name", "field_type",
                   "src_file", "src_type", "src_locator", "description"]
EDGE_COLUMNS = ["from_dataset", "from_field", "to_dataset", "to_field",
                "transform", "confidence", "note"]


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _parse_locator(raw, row: int) -> dict:
    """src_locator 接受 dict（JSON 清单）或 JSON 字符串（CSV/XLSX 单元格）。"""
    if raw is None:
        raise ValueError(f"第 {row} 行：src_locator 缺失")
    if isinstance(raw, dict):
        return raw
    try:
        v = json.loads(str(raw))
    except json.JSONDecodeError as e:
        raise ValueError(f"第 {row} 行：src_locator 不是合法 JSON（{e.msg}）") from None
    if not isinstance(v, dict):
        raise ValueError(f"第 {row} 行：src_locator 必须是 JSON 对象")
    return v


def _row_record(kind: str, data: dict, row: int):
    """清单行 → 记录。行级解析错误不抛出，写入 parse_error 由导入阶段整行拒绝；
    文件级错误（缺 sheet、根解析失败）仍由调用方整批 400。"""
    if kind == "field":
        base = {
            "row": row,
            "dataset_id": _clean(data.get("dataset_id")) or "",
            "dataset_name": _clean(data.get("dataset_name")),
            "field_name": _clean(data.get("field_name")) or "",
            "field_type": _clean(data.get("field_type")),
            "src_file": _clean(data.get("src_file")) or "",
            "src_type": (_clean(data.get("src_type")) or "").lower(),
            "description": _clean(data.get("description")),
        }
        try:
            base["src_locator"] = _parse_locator(data.get("src_locator"), row)
        except ValueError as e:
            # 解析失败的行注定被拒，src_locator 以空对象占位
            base["src_locator"] = {}
            base["parse_error"] = str(e)
        return DatasetFieldRecord(**base)
    base = {
        "row": row,
        "from_dataset": _clean(data.get("from_dataset")) or "",
        "from_field": _clean(data.get("from_field")),
        "to_dataset": _clean(data.get("to_dataset")) or "",
        "to_field": _clean(data.get("to_field")),
        "transform": _clean(data.get("transform")),
        "note": _clean(data.get("note")),
    }
    confidence = _clean(data.get("confidence"))
    try:
        base["confidence"] = float(confidence) if confidence else None
    except ValueError:
        base["confidence"] = None
        base["parse_error"] = f"第 {row} 行：confidence 不是合法数字"
    return LineageEdgeRecord(**base)


def _headers_to_dicts(header_row, rows) -> list[dict]:
    headers = [_clean(h) for h in header_row]
    return [dict(zip(headers, r, strict=False)) for r in rows]


def read_csv_manifest(dataset_bytes: bytes, lineage_bytes: bytes
                      ) -> tuple[list[DatasetFieldRecord], list[LineageEdgeRecord]]:
    ds_recs, edge_recs = [], []
    reader = csv.reader(io.StringIO(dataset_bytes.decode("utf-8-sig")))
    rows = list(reader)
    for i, data in enumerate(_headers_to_dicts(rows[0], rows[1:]), start=2):
        ds_recs.append(_row_record("field", data, i))
    reader = csv.reader(io.StringIO(lineage_bytes.decode("utf-8-sig")))
    rows = list(reader)
    for i, data in enumerate(_headers_to_dicts(rows[0], rows[1:]), start=2):
        edge_recs.append(_row_record("edge", data, i))
    return ds_recs, edge_recs


def read_xlsx_manifest(data: bytes) -> tuple[list[DatasetFieldRecord], list[LineageEdgeRecord]]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    by_name = {n.strip().lower(): wb[n] for n in wb.sheetnames}
    for required in ("datasets", "lineage"):
        if required not in by_name:
            raise ValueError(f"XLSX 清单缺少工作表「{required}」（现有：{', '.join(wb.sheetnames)}）")

    def sheet_dicts(name: str) -> list[dict]:
        ws = by_name[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        return _headers_to_dicts(rows[0], rows[1:])

    ds_recs = [_row_record("field", d, i) for i, d in enumerate(sheet_dicts("datasets"), start=2)]
    edge_recs = [_row_record("edge", d, i) for i, d in enumerate(sheet_dicts("lineage"), start=2)]
    wb.close()
    return ds_recs, edge_recs


def read_json_manifest(data: bytes
                       ) -> tuple[list[DatasetFieldRecord], list[LineageEdgeRecord]]:
    try:
        doc = json.loads(data.decode("utf-8-sig"))
    except json.JSONDecodeError as e:
        raise ValueError(f"JSON 清单解析失败（{e.msg}）") from None
    ds_recs = [_row_record("field", d, i) for i, d in enumerate(doc.get("datasets", []), start=1)]
    edge_recs = [_row_record("edge", d, i) for i, d in enumerate(doc.get("lineage", []), start=1)]
    return ds_recs, edge_recs
