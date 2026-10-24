"""源文件：构造含多级表头 + 合并单元格的 XLSX、CSV、JSON 第一手文件。"""

import io
import json

from openpyxl import Workbook


def order_xlsx_bytes() -> bytes:
    """订单明细.xlsx：第 1 行 B1:F1 合并「商品信息」、G1「运费信息」；
    第 2 行多级表头；第 3-4 行数据。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["B1"] = "商品信息"
    ws.merge_cells("B1:F1")
    ws["G1"] = "运费信息"
    for col, name in zip("BCDEFG", ["商品ID", "订单号", "数量", "单价", "金额", "运费"],
                         strict=True):
        ws[f"{col}2"] = name
    rows = [("S001", "SO-1001", 2, 9.9, 19.8, 5),
            ("S002", "SO-1002", 1, 59.0, 59.0, 0)]
    for i, row in enumerate(rows, start=3):
        for j, v in enumerate(row, start=2):
            ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def user_csv_bytes() -> bytes:
    return ("user_id,user_name,phone\n"
            "1,张三,138*\n"
            "2,李四,139*\n").encode("utf-8")


def source_json_bytes() -> bytes:
    return json.dumps({"data": {"items": [{"order_no": "SO-1001"},
                                          {"order_no": "SO-1002"}]}},
                      ensure_ascii=False).encode("utf-8")
