"""验收：三种格式导入 + 定位器实文件核查（含 XLSX 合并单元格/多级表头）。"""

import io
import json

from openpyxl import Workbook

from tests.conftest import csv_bytes, login, upload
from tests.fixtures import order_xlsx_bytes, source_json_bytes, user_csv_bytes


def _seed_sources(client, headers) -> None:
    upload(client, headers, [
        ("订单明细.xlsx", order_xlsx_bytes(),
         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("用户表.csv", user_csv_bytes(), "text/csv"),
        ("源.json", source_json_bytes(), "application/json"),
    ])


def _import_csv(client, headers, ds_rows, edge_rows):
    return client.post("/import", headers=headers,
                       data={"format": "csv", "source_file_ids": "1,2,3"},
                       files={
                           "manifest_datasets": ("datasets.csv", csv_bytes(ds_rows), "text/csv"),
                           "manifest_lineage": ("lineage.csv", csv_bytes(edge_rows), "text/csv"),
                       })


def test_csv_import_with_multiheader_xlsx_and_bad_locator(client):
    headers = login(client)
    _seed_sources(client, headers)

    ds_rows = [
        ["dataset_id", "dataset_name", "field_name", "field_type", "src_file",
         "src_type", "src_locator", "description"],
        ["ods_raw", "订单明细", "order_id", "STRING", "订单明细.xlsx", "xlsx",
         json.dumps({"sheet": "订单", "header_path": ["商品信息", "商品ID"],
                     "header_rows": 2, "row_start": 3}), "商品ID"],
        ["ods_raw", "订单明细", "order_no", "STRING", "订单明细.xlsx", "xlsx",
         json.dumps({"sheet": "订单", "header_path": ["商品信息", "订单号"],
                     "header_rows": 2, "row_start": 3}), "订单号"],
        ["ods_raw", "订单明细", "amount", "DECIMAL", "订单明细.xlsx", "xlsx",
         json.dumps({"sheet": "订单", "header_path": ["商品信息", "金额"],
                     "header_rows": 2, "row_start": 3}), "金额"],
        ["ods_raw", "订单明细", "bad_field", "STRING", "订单明细.xlsx", "xlsx",
         json.dumps({"sheet": "订单", "header_path": ["商品信息", "不存在"], "header_rows": 2}),
         "错误定位器"],
        ["dwd_order", "订单宽表", "order_no", "STRING", "订单明细.xlsx", "xlsx",
         json.dumps({"sheet": "订单", "column": "C", "row_start": 3}), "订单号宽表"],
        ["dwd_order", "订单宽表", "user_name", "STRING", "用户表.csv", "csv",
         json.dumps({"column": 2}), "用户名"],
    ]
    edge_rows = [
        ["from_dataset", "from_field", "to_dataset", "to_field", "transform",
         "confidence", "note"],
        ["ods_raw", "order_id", "dwd_order", "order_no", "trim/cast", "1.0", "同源加工"],
        ["ods_raw", "", "dwd_order", "", "", "", ""],
    ]
    r = _import_csv(client, headers, ds_rows, edge_rows)
    assert r.status_code == 201, r.text
    report = r.json()["report"]
    assert report["summary"] == {"total": 8, "ok": 7, "merged": 0, "rejected": 1}
    rejected = [i for i in report["items"] if i["status"] == "rejected"]
    assert rejected[0]["ref"] == "ods_raw.bad_field"
    assert "实文件核查失败" in rejected[0]["reason"]

    # 错误定位器的字段不入库；正确字段带物化解析结果
    ods = client.get("/datasets/ods_raw", headers=headers).json()
    assert [f["name"] for f in ods["fields"]] == ["amount", "order_id", "order_no"]
    order_no = next(f for f in ods["fields"] if f["name"] == "order_no")
    resolved = order_no["src_locator"]["resolved"]
    assert resolved["column_index"] == 3  # 订单号在 C 列
    assert resolved["merged_anchors"] == [
        {"value": "商品信息", "anchor": "B1", "range": "B1:F1"}]

    dwd = client.get("/datasets/dwd_order", headers=headers).json()
    simple = next(f for f in dwd["fields"] if f["name"] == "order_no")
    assert simple["src_locator"]["input"] == {"sheet": "订单", "column": "C", "row_start": 3}
    assert simple["src_locator"]["resolved"]["column_index"] == 3


def test_xlsx_manifest_import(client):
    headers = login(client)
    _seed_sources(client, headers)

    wb = Workbook()
    ds = wb.active
    ds.title = "datasets"
    ds.append(["dataset_id", "dataset_name", "field_name", "field_type",
               "src_file", "src_type", "src_locator", "description"])
    ds.append(["ods_raw", "订单明细", "order_no", "STRING", "订单明细.xlsx", "xlsx",
               json.dumps({"sheet": "订单", "header_path": ["商品信息", "订单号"],
                           "header_rows": 2, "row_start": 3}), "订单号"])
    ds.append(["ods_raw", "订单明细", "qty", "INTEGER", "订单明细.xlsx", "xlsx",
               json.dumps({"sheet": "订单", "column": "D", "row_start": 3}), "数量"])
    ds.append(["ods_raw", "订单明细", "broken", "STRING", "订单明细.xlsx", "xlsx",
               json.dumps({"sheet": "不存在的工作表", "column": "A"}), "错 sheet"])
    lineage = wb.create_sheet("lineage")
    lineage.append(["from_dataset", "from_field", "to_dataset", "to_field",
                    "transform", "confidence", "note"])
    lineage.append(["ods_raw", "order_no", "ods_raw", "qty", "derive", "1.0", ""])
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post("/import", headers=headers,
                    data={"format": "xlsx", "source_file_ids": "1,2,3"},
                    files={"manifest": ("manifest.xlsx", buf.getvalue(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 201, r.text
    report = r.json()["report"]
    assert report["summary"] == {"total": 4, "ok": 3, "merged": 0, "rejected": 1}

    ods = client.get("/datasets/ods_raw", headers=headers).json()
    assert [f["name"] for f in ods["fields"]] == ["order_no", "qty"]
    assert ods["fields"][1]["src_locator"]["resolved"]["column_index"] == 4


def test_json_manifest_import(client):
    headers = login(client)
    _seed_sources(client, headers)

    manifest = {
        "datasets": [
            {"dataset_id": "ods_user", "dataset_name": "用户表", "field_name": "user_id",
             "field_type": "INTEGER", "src_file": "用户表.csv", "src_type": "csv",
             "src_locator": {"header": "user_id"}},
            {"dataset_id": "ods_user", "dataset_name": "用户表", "field_name": "user_name",
             "field_type": "STRING", "src_file": "用户表.csv", "src_type": "csv",
             "src_locator": {"column": 2}},
            {"dataset_id": "ods_user", "dataset_name": "用户表", "field_name": "ext_order_no",
             "field_type": "STRING", "src_file": "源.json", "src_type": "json",
             "src_locator": {"json_path": "$.data.items[*].order_no"}},
        ],
        "lineage": [
            {"from_dataset": "ods_user", "from_field": "user_id",
             "to_dataset": "ods_user", "to_field": "user_name", "transform": "lookup"},
        ],
    }
    r = client.post("/import", headers=headers,
                    data={"format": "json", "source_file_ids": "1,2,3"},
                    files={"manifest": ("manifest.json", json.dumps(manifest, ensure_ascii=False).encode(),
                                        "application/json")})
    assert r.status_code == 201, r.text
    report = r.json()["report"]
    assert report["summary"] == {"total": 4, "ok": 4, "merged": 0, "rejected": 0}

    ods = client.get("/datasets/ods_user", headers=headers).json()
    ext = next(f for f in ods["fields"] if f["name"] == "ext_order_no")
    assert ext["src_locator"]["resolved"]["hits"] == 2  # JSONPath 命中数


def test_manifest_missing_sheet_rejected(client):
    headers = login(client)
    _seed_sources(client, headers)
    wb = Workbook()
    wb.active.title = "wrong"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/import", headers=headers,
                    data={"format": "xlsx", "source_file_ids": "1"},
                    files={"manifest": ("m.xlsx", buf.getvalue(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 400
    assert "datasets" in r.json()["detail"]
