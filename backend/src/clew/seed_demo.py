"""M2 演示数据：5 个数据集、4 跳血缘链（设计方案 §9 M2 验收样例）。

执行方式::

    cd backend && uv run python -m clew.seed_demo

现场生成 csv / xlsx（含合并单元格多级表头）/ json 三类源文件，先走上传
服务（sha256 去重）登记，再走 ``read_json_manifest`` + ``run_import``
正式导入路径入库。重复执行时检测到完整演示数据即跳过，保证幂等。
"""

import csv as csvmod
import io
import json
import sys
from datetime import UTC, datetime

from openpyxl import Workbook

from .config import get_settings
from .db import get_db, init_db
from .import_service import run_import
from .manifest import read_json_manifest
from .routes.files import store_upload
from .security import hash_password

DEMO_DATASET_IDS = ("d0", "d1", "d2", "d3", "d4")

# (dataset_id, dataset_name, [(field_name, field_type, src_file, src_type,
#                              src_locator, description)])
DEMO_DATASETS = [
    ("d0", "原始订单", [
        ("order_id", "STRING", "d0_orders.csv", "csv", {"header": "order_id"}, "订单号"),
        ("amount", "DECIMAL", "d0_orders.csv", "csv", {"header": "amount"}, "订单金额"),
        ("region", "STRING", "d0_orders.csv", "csv", {"header": "region"}, "销售地区"),
    ]),
    ("d1", "订单清洗", [
        ("order_id", "STRING", "d1_orders.xlsx", "xlsx",
         {"sheet": "订单", "header_path": ["订单信息", "订单号"],
          "header_rows": 2, "row_start": 3}, "清洗后订单号"),
        ("amount", "DECIMAL", "d1_orders.xlsx", "xlsx",
         {"sheet": "订单", "header_path": ["订单信息", "金额"],
          "header_rows": 2, "row_start": 3}, "清洗后金额"),
        ("region", "STRING", "d1_orders.xlsx", "xlsx",
         {"sheet": "订单", "header_path": ["订单信息", "地区"],
          "header_rows": 2, "row_start": 3}, "标准化地区"),
    ]),
    ("d2", "订单明细", [
        ("order_id", "STRING", "d2_orders.json", "json",
         {"json_path": "$.data.orders[*].order_id"}, "明细订单号"),
        ("amount", "DECIMAL", "d2_orders.json", "json",
         {"json_path": "$.data.orders[*].amount"}, "明细金额"),
        ("region", "STRING", "d2_orders.json", "json",
         {"json_path": "$.data.orders[*].region"}, "明细地区"),
    ]),
    ("d3", "日汇总", [
        ("order_id", "STRING", "d3_daily.csv", "csv", {"header": "order_id"}, "汇总订单号"),
        ("amount", "DECIMAL", "d3_daily.csv", "csv", {"header": "amount"}, "日汇总金额"),
        ("region", "STRING", "d3_daily.csv", "csv", {"header": "region"}, "汇总地区"),
    ]),
    ("d4", "报表", [
        ("order_id", "STRING", "d4_report.json", "json",
         {"json_path": "$.report.rows[*].order_id"}, "报表订单号"),
        ("amount", "DECIMAL", "d4_report.json", "json",
         {"json_path": "$.report.rows[*].amount"}, "报表金额"),
        ("region", "STRING", "d4_report.json", "json",
         {"json_path": "$.report.rows[*].region"}, "报表地区"),
    ]),
]

_ROWS = [
    ("SO-1001", 19.8, "华东"),
    ("SO-1002", 59.0, "华北"),
    ("SO-1003", 12.5, "华南"),
]


def _csv_bytes(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csvmod.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8")


def _d0_csv() -> bytes:
    return _csv_bytes([["order_id", "amount", "region"], *[list(r) for r in _ROWS]])


def _d3_csv() -> bytes:
    # 与 d0 内容不同（日汇总额），避免 sha256 去重后无法按文件名绑定
    rows = [[order_id, round(amount * 2, 2), region]
            for order_id, amount, region in _ROWS]
    return _csv_bytes([["order_id", "amount", "region"], *rows])


def _d1_xlsx() -> bytes:
    """订单清洗.xlsx：B1:F1 合并为多级表头「订单信息」。"""
    wb = Workbook()
    # 固定文档时间戳，尽量让重复生成的字节稳定（幂等辅助；主保障是已完成检测）
    wb.properties.created = datetime(2026, 1, 1)
    wb.properties.modified = datetime(2026, 1, 1)
    ws = wb.active
    ws.title = "订单"
    ws["B1"] = "订单信息"
    ws.merge_cells("B1:F1")
    for col, name in zip("BCDEF", ["订单号", "金额", "地区", "备注", "来源"],
                         strict=True):
        ws[f"{col}2"] = name
    for row_idx, (order_id, amount, region) in enumerate(_ROWS, start=3):
        ws.cell(row_idx, 2, order_id)
        ws.cell(row_idx, 3, amount)
        ws.cell(row_idx, 4, region)
        ws.cell(row_idx, 5, "演示")
        ws.cell(row_idx, 6, "seed_demo")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _d2_json() -> bytes:
    return json.dumps(
        {"data": {"orders": [
            {"order_id": order_id, "amount": amount, "region": region}
            for order_id, amount, region in _ROWS]}},
        ensure_ascii=False).encode("utf-8")


def _d4_json() -> bytes:
    return json.dumps(
        {"report": {"rows": [
            {"order_id": order_id, "amount": amount, "region": region}
            for order_id, amount, region in _ROWS]}},
        ensure_ascii=False).encode("utf-8")


def _demo_files() -> dict[str, bytes]:
    return {
        "d0_orders.csv": _d0_csv(),
        "d1_orders.xlsx": _d1_xlsx(),
        "d2_orders.json": _d2_json(),
        "d3_daily.csv": _d3_csv(),
        "d4_report.json": _d4_json(),
    }


def _manifest(file_ids: dict[str, int]) -> dict:
    datasets = []
    for dataset_id, dataset_name, fields in DEMO_DATASETS:
        for field_name, field_type, src_file, src_type, locator, description in fields:
            datasets.append({
                "dataset_id": dataset_id,
                "dataset_name": dataset_name,
                "field_name": field_name,
                "field_type": field_type,
                "src_file": str(file_ids[src_file]),
                "src_type": src_type,
                "src_locator": locator,
                "description": description,
            })
    lineage = []
    for i in range(len(DEMO_DATASET_IDS) - 1):
        src, dst = DEMO_DATASET_IDS[i], DEMO_DATASET_IDS[i + 1]
        lineage.append({
            "from_dataset": src, "from_field": "order_id",
            "to_dataset": dst, "to_field": "order_id",
            "transform": "trim + cast + 标准化",
            "confidence": 1.0, "note": "M2 演示：字段级血缘",
        })
        lineage.append({
            "from_dataset": src, "from_field": None,
            "to_dataset": dst, "to_field": None,
            "transform": "表级投影",
            "confidence": 1.0, "note": "M2 演示：表级血缘",
        })
    return {"datasets": datasets, "lineage": lineage}


def _already_seeded(conn) -> bool:
    """演示数据集、全部演示字段与全部演示边均已存在时视为已完成。"""
    marks = ",".join("?" for _ in DEMO_DATASET_IDS)
    dataset_ids = {r["id"] for r in conn.execute(
        f"SELECT id FROM dataset WHERE id IN ({marks})", DEMO_DATASET_IDS)}
    if dataset_ids != set(DEMO_DATASET_IDS):
        return False

    expected_fields = {
        (dataset_id, field[0])
        for dataset_id, _, fields in DEMO_DATASETS
        for field in fields
    }
    actual_fields = {
        (r["dataset_id"], r["name"]) for r in conn.execute(
            f"SELECT dataset_id, name FROM field WHERE dataset_id IN ({marks})",
            DEMO_DATASET_IDS)
    }
    if not expected_fields <= actual_fields:
        return False

    expected_edges = set()
    for i in range(len(DEMO_DATASET_IDS) - 1):
        src, dst = DEMO_DATASET_IDS[i], DEMO_DATASET_IDS[i + 1]
        expected_edges.add((src, "order_id", dst, "order_id"))
        expected_edges.add((src, None, dst, None))
    actual_edges = {
        (r["from_dataset"], r["from_field"], r["to_dataset"], r["to_field"])
        for r in conn.execute(
            f"SELECT from_dataset, from_field, to_dataset, to_field FROM edge"
            f" WHERE from_dataset IN ({marks}) AND to_dataset IN ({marks})",
            [*DEMO_DATASET_IDS, *DEMO_DATASET_IDS])
    }
    return expected_edges <= actual_edges


def _ensure_admin() -> str:
    """演示数据需要一个合法 actor（source_file.uploaded_by 外键）。"""
    settings = get_settings()
    with get_db() as conn:
        exists = conn.execute("SELECT 1 FROM user WHERE username=?",
                              (settings.admin_username,)).fetchone()
        if exists is None:
            conn.execute(
                "INSERT INTO user(username, password_hash, display_name, role, created_at)"
                " VALUES (?,?,?,?,?)",
                (settings.admin_username, hash_password(settings.admin_password),
                 "初始管理员", "admin",
                 datetime.now(UTC).isoformat(timespec="seconds")))
    return settings.admin_username


def seed_demo() -> dict:
    """构造演示数据；返回本次执行摘要（已存在时 skipped=True）。"""
    init_db()
    actor = _ensure_admin()

    with get_db() as conn:
        if _already_seeded(conn):
            print("演示数据已完整存在，跳过（幂等）", file=sys.stderr)
            return {"skipped": True, "batch_id": None, "version_no": None,
                    "summary": None}

    bound, file_ids = [], {}
    for name, data in _demo_files().items():
        rec = store_upload(name, data, actor)
        bound.append(rec)
        file_ids[name] = rec["id"]
    manifest_bytes = json.dumps(_manifest(file_ids), ensure_ascii=False).encode()
    ds_recs, edge_recs = read_json_manifest(manifest_bytes)
    batch_id, version_no, report = run_import(actor, "json", bound,
                                              ds_recs, edge_recs)
    print(f"演示数据导入完成：批次 v{version_no}，"
          f"数据集 {'→'.join(DEMO_DATASET_IDS)}，报告 {report.summary}",
          file=sys.stderr)
    return {"skipped": False, "batch_id": batch_id, "version_no": version_no,
            "summary": report.summary}


if __name__ == "__main__":
    seed_demo()
