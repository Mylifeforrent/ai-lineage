"""字段回源预览：按 ``field.src_locator`` 信封打开托管源文件并返回高亮所需数据。

返回结构与 API 契约一致（``csv`` / ``xlsx`` / ``json`` 三选一）：

- csv：首列表头 + 列号 + 前 50 数据行；
- xlsx：sheet、表头行、合并区域、列号、数据行范围 + 前 50 数据行；
- json：JSONPath、命中总数 + 前 10 个命中值。

定位器已失效（文件缺失、sheet 不存在、物化行列范围越界、JSONPath 无命中等）统一抛
:class:`PreviewError`，由路由层转 409；数据集/字段不存在由调用方转 404。
"""

import csv
import json
import sqlite3
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from jsonpath_ng import parse as jsonpath_parse
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .config import resolve_storage

PREVIEW_ROW_LIMIT = 50
PREVIEW_JSON_LIMIT = 10


class PreviewError(ValueError):
    """定位器对当前源文件已失效（→ 409）。"""


def _cell_value(value):
    """openpyxl 单元格值 → JSON 可序列化值。"""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _envelope(raw: str | None) -> dict:
    try:
        value = json.loads(raw or "")
    except json.JSONDecodeError:
        raise PreviewError("字段定位器不是合法 JSON，定位已失效") from None
    if not isinstance(value, dict):
        raise PreviewError("字段定位器结构不合法，定位已失效")
    return value


def _preview_csv(path: Path, locator: dict) -> dict:
    resolved = locator.get("resolved") or {}
    user_input = locator.get("input") or {}

    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        if header is None:
            raise PreviewError("CSV 源文件为空，无法预览")
        rows = [row for _, row in zip(range(PREVIEW_ROW_LIMIT), reader, strict=False)]

    column_index = resolved.get("column_index")
    if column_index is None:
        if user_input.get("header") is not None:
            try:
                column_index = header.index(user_input["header"]) + 1
            except ValueError:
                raise PreviewError(
                    f"CSV 表头「{user_input['header']}」未在源文件中找到，定位已失效") from None
        elif user_input.get("column") is not None:
            column_index = user_input["column"]
    if not isinstance(column_index, int) or not 1 <= column_index <= len(header):
        raise PreviewError(
            f"CSV 列定位（{column_index}）超出源文件列数（{len(header)}），定位已失效")

    width = len(header)
    rows = [row + [""] * (width - len(row)) if len(row) < width else row for row in rows]
    return {"header": header, "column_index": column_index, "rows": rows}


def _preview_xlsx(path: Path, locator: dict) -> dict:
    resolved = locator.get("resolved") or {}
    user_input = locator.get("input") or {}
    sheet = user_input.get("sheet")
    if not sheet:
        raise PreviewError("XLSX 定位器缺少 sheet，定位已失效")

    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        if sheet not in wb.sheetnames:
            raise PreviewError(
                f"工作表「{sheet}」已不存在（现有：{', '.join(wb.sheetnames)}），定位已失效")
        ws = wb[sheet]
        max_col = ws.max_column
        max_row = ws.max_row

        column_index = resolved.get("column_index")
        if column_index is None:
            column = user_input.get("column") or resolved.get("column")
            if not column:
                raise PreviewError("XLSX 定位器缺少列信息，定位已失效")
            try:
                column_index = column_index_from_string(str(column).upper())
            except ValueError:
                raise PreviewError(
                    f"XLSX 列标识「{column}」不合法，定位已失效") from None
        if not isinstance(column_index, int) or not 1 <= column_index <= max_col:
            raise PreviewError(
                f"XLSX 列定位（{column_index}）超出工作表列数（{max_col}），定位已失效")

        header_count = user_input.get("header_rows")
        if header_count is None and user_input.get("header_path"):
            header_count = len(user_input["header_path"])
        header_count = int(header_count or 0)
        header_count = max(0, min(header_count, max_row))
        header_rows = [
            [_cell_value(ws.cell(r, c).value) for c in range(1, max_col + 1)]
            for r in range(1, header_count + 1)
        ]

        declared_start = resolved.get("row_start")
        if declared_start is None:
            declared_start = user_input.get("row_start")
        if declared_start is None:
            # 未物化行范围：默认从表头后第一行开始（空数据表仍返回空预览）
            row_start = header_count + 1
        else:
            row_start = int(declared_start)
            if row_start > max_row:
                raise PreviewError(
                    f"数据行起点（row_start={row_start}）超出工作表行数"
                    f"（{max_row}），定位已失效")

        declared_end = resolved.get("row_end")
        if declared_end is None:
            declared_end = user_input.get("row_end")
        if declared_end is None:
            row_end = max_row
        else:
            row_end = int(declared_end)
            if row_end > max_row:
                raise PreviewError(
                    f"数据行终点（row_end={row_end}）超出工作表行数"
                    f"（{max_row}），定位已失效")
        if row_end < row_start - 1:
            raise PreviewError(
                f"数据行范围无效（row_end={row_end} 小于 row_start={row_start}），定位已失效")
        preview_end = min(row_end, row_start + PREVIEW_ROW_LIMIT - 1)
        rows = []
        if preview_end >= row_start:
            rows = [
                [_cell_value(ws.cell(r, c).value) for c in range(1, max_col + 1)]
                for r in range(row_start, preview_end + 1)
            ]

        # 合并区域：返回与「表头 + 本次预览行」相交的区域，供前端渲染合并单元格
        visible_end = max(header_count, preview_end, 0)
        merged_ranges = [
            str(rng) for rng in sorted(ws.merged_cells.ranges,
                                       key=lambda r: (r.min_row, r.min_col))
            if rng.min_row <= visible_end
        ]
        return {
            "sheet": sheet,
            "header_rows": header_rows,
            "merged_ranges": merged_ranges,
            "column_index": column_index,
            "row_start": row_start,
            "row_end": row_end if row_end >= row_start else row_start - 1,
            "rows": rows,
        }
    finally:
        wb.close()


def _preview_json(path: Path, locator: dict) -> dict:
    user_input = locator.get("input") or {}
    json_path = user_input.get("json_path")
    if not json_path:
        raise PreviewError("JSON 定位器缺少 json_path，定位已失效")

    try:
        with open(path, encoding="utf-8-sig") as f:
            doc = json.load(f)
    except json.JSONDecodeError as e:
        raise PreviewError(f"源 JSON 文件解析失败（{e.msg}），定位已失效") from None

    try:
        expr = jsonpath_parse(json_path)
    except Exception:
        raise PreviewError(f"JSONPath「{json_path}」语法不合法，定位已失效") from None

    matches = expr.find(doc)
    if not matches:
        raise PreviewError(f"JSONPath「{json_path}」在源文件中无命中，定位已失效")
    return {
        "json_path": json_path,
        "hits": len(matches),
        "samples": [m.value for m in matches[:PREVIEW_JSON_LIMIT]],
    }


_PREVIEWERS = {"csv": _preview_csv, "xlsx": _preview_xlsx, "json": _preview_json}


def field_preview(conn: sqlite3.Connection, dataset_id: str,
                  field_name: str) -> dict:
    """返回字段的回源预览；字段/数据集不存在抛 LookupError。"""
    if conn.execute("SELECT 1 FROM dataset WHERE id=?", (dataset_id,)).fetchone() is None:
        raise LookupError(f"数据集「{dataset_id}」不存在")
    row = conn.execute(
        "SELECT f.name, f.src_locator, s.id AS source_file_id,"
        " s.original_name, s.storage_path, s.src_type, s.sha256,"
        " s.size_bytes, s.uploaded_by"
        " FROM field f JOIN source_file s ON s.id = f.source_file_id"
        " WHERE f.dataset_id=? AND f.name=?",
        (dataset_id, field_name)).fetchone()
    if row is None:
        raise LookupError(f"字段「{dataset_id}.{field_name}」不存在")

    locator = _envelope(row["src_locator"])
    src_type = locator.get("type") or row["src_type"]
    if src_type != row["src_type"]:
        raise PreviewError(
            f"字段定位器类型「{src_type}」与源文件类型「{row['src_type']}」不一致，定位已失效")
    previewer = _PREVIEWERS.get(src_type)
    if previewer is None:
        raise PreviewError(f"不支持的源文件类型「{src_type}」")

    path = resolve_storage(row["storage_path"])
    if not path.is_file():
        raise PreviewError(
            f"源文件「{row['original_name']}」不在托管存储中，定位已失效")
    try:
        preview = previewer(path, locator)
    except PreviewError:
        raise
    except Exception as e:
        # 损坏的 csv/xlsx/json 可能让解析器抛多种异常，统一按「定位失效」处理
        raise PreviewError(
            f"源文件「{row['original_name']}」解析失败（{e}），定位已失效") from None

    # 实文件核查记录：字段最近一次入库（added/updated）的导入批次。
    # 导入器只有通过实文件核查才会写 change_log，因此该记录即「核查通过」。
    verified = conn.execute(
        "SELECT b.version_no, cl.actor, cl.created_at"
        " FROM change_log cl JOIN import_batch b ON b.id = cl.batch_id"
        " WHERE cl.entity_kind='field' AND cl.entity_key=?"
        " AND cl.op IN ('added','updated') ORDER BY cl.id DESC LIMIT 1",
        (f"{dataset_id}.{field_name}",)).fetchone()

    return {
        "file": {
            "id": row["source_file_id"],
            "original_name": row["original_name"],
            # 源文件按 sha256 内容寻址去重，前端展示托管/去重依据
            "sha256": row["sha256"],
            "size_bytes": row["size_bytes"],
            "src_type": row["src_type"],
            "uploaded_by": row["uploaded_by"],
        },
        "verification": ({
            "status": "passed",
            "version_no": verified["version_no"],
            "checked_by": verified["actor"],
            "checked_at": verified["created_at"],
        } if verified is not None else None),
        "locator": locator,
        "preview": preview,
    }
