"""定位器实文件核查（设计方案 §3.2 / §3.5-2）。

导入时打开托管的第一手文件机械核查定位器，并物化解析结果：
- csv：列号/表头解析为 column_index
- xlsx：sheet 存在性、简单列定位、多级表头 header_path 展开、合并单元格锚点归一
- json：JSONPath 命中数
"""

import csv
import io
import json

from jsonpath_ng import parse as jsonpath_parse
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .models import CsvLocatorIn, JsonLocatorIn, XlsxLocatorIn


class LocatorError(ValueError):
    """定位器核查失败（结构校验或实文件核查不通过）。"""


def check_csv(path: str, loc: CsvLocatorIn) -> dict:
    with open(path, newline="", encoding="utf-8-sig") as f:
        header = next(csv.reader(f), [])
    width = len(header)
    if loc.header is not None:
        for i, name in enumerate(header, start=1):
            if name and name.strip() == loc.header:
                return {"column_index": i, "header": loc.header}
        raise LocatorError(f"csv 表头「{loc.header}」未在源文件首行找到（共 {width} 列）")
    if loc.column is None:
        raise LocatorError("csv 定位器必须提供 column 或 header 之一")
    if loc.column > width:
        raise LocatorError(f"csv 列号 {loc.column} 超出源文件列数（{width}）")
    return {"column_index": loc.column,
            "header": header[loc.column - 1] if loc.column <= width else None}


def _merged_anchor(ws, row_idx: int, col_idx: int) -> dict | None:
    """若 (row, col) 落在合并区域内，返回锚点信息（左上角 + 范围）。"""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row_idx <= rng.max_row and rng.min_col <= col_idx <= rng.max_col:
            return {
                "value": ws.cell(rng.min_row, rng.min_col).value,
                "anchor": f"{get_column_letter(rng.min_col)}{rng.min_row}",
                "range": rng.coord,
            }
    return None


def check_xlsx(path: str, loc: XlsxLocatorIn) -> dict:
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        if loc.sheet not in wb.sheetnames:
            raise LocatorError(
                f"xlsx sheet「{loc.sheet}」不存在（现有：{', '.join(wb.sheetnames)}）")
        ws = wb[loc.sheet]

        if loc.header_path:
            header_rows = loc.header_rows or len(loc.header_path)
            resolved, candidates, anchors = None, None, []
            for level, want in enumerate(loc.header_path, start=1):
                if level > header_rows:
                    raise LocatorError(f"表头路径层级 {level} 超出 header_rows={header_rows}")
                hit_col = None
                for col in range(1, ws.max_column + 1):
                    if candidates and col not in candidates:
                        continue
                    v = ws.cell(level, col).value
                    if v is not None and str(v).strip() == want:
                        hit_col = col
                        break
                if hit_col is None:
                    scope = f"列 {get_column_letter(min(candidates))}-{get_column_letter(max(candidates))}" \
                        if candidates else f"第 {level} 行全表"
                    raise LocatorError(
                        f"表头路径第 {level} 层「{want}」未匹配到（在 {scope} 内查找）")
                anchor = _merged_anchor(ws, level, hit_col)
                if anchor:
                    anchors.append(anchor)
                    rng = next(r for r in ws.merged_cells.ranges if r.coord == anchor["range"])
                    candidates = set(range(rng.min_col, rng.max_col + 1))
                else:
                    candidates = {hit_col}
                resolved = hit_col
            row_start = loc.row_start or (header_rows + 1)
            return {
                "column_index": resolved,
                "column": get_column_letter(resolved),
                "row_start": row_start,
                "row_end": loc.row_end,
                "merged_anchors": anchors,
            }

        if not loc.column:
            raise LocatorError("xlsx 定位器必须提供 column 或 header_path 之一")
        try:
            col_idx = column_index_from_string(loc.column.upper())
        except ValueError:
            raise LocatorError(f"xlsx 列标识「{loc.column}」不合法") from None
        if loc.row_start and loc.row_end and loc.row_end < loc.row_start:
            raise LocatorError("row_end 不能小于 row_start")
        return {"column_index": col_idx, "column": loc.column.upper(),
                "row_start": loc.row_start, "row_end": loc.row_end, "merged_anchors": []}
    finally:
        wb.close()


def check_json(path: str, loc: JsonLocatorIn) -> dict:
    with open(path, encoding="utf-8-sig") as f:
        try:
            doc = json.load(f)
        except json.JSONDecodeError as e:
            raise LocatorError(f"源 JSON 文件解析失败（{e.msg}）") from None
    try:
        expr = jsonpath_parse(loc.json_path)
    except Exception:
        raise LocatorError(f"JSONPath「{loc.json_path}」语法不合法") from None
    hits = len(expr.find(doc))
    if hits == 0:
        raise LocatorError(f"JSONPath「{loc.json_path}」在源文件中无命中")
    return {"hits": hits}


# src_type → 结构模型/实文件核查器；导入与 M3 字段定位器编辑共用
CHECKERS = {"csv": check_csv, "xlsx": check_xlsx, "json": check_json}
LOCATOR_MODELS = {"csv": CsvLocatorIn, "xlsx": XlsxLocatorIn,
                  "json": JsonLocatorIn}
